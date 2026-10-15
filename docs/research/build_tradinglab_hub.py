import json
import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.stdout.reconfigure(encoding='utf-8')
research_dir = r"D:\victor\Trading\docs\research"
results_file = os.path.join(research_dir, "tradinglab_results.json")

if not os.path.exists(results_file):
    print("tradinglab_results.json chưa sẵn sàng.")
    sys.exit(0)

with open(results_file, "r", encoding="utf-8") as f:
    videos = json.load(f)

print(f"Bắt đầu tổng hợp {len(videos)} video từ TradingLab...")

# 1. Update MASTER_VIDEO_DATABASE.xlsx with a new dedicated sheet
wb_path = os.path.join(research_dir, "MASTER_VIDEO_DATABASE.xlsx")
wb = openpyxl.load_workbook(wb_path)

sheet_name = "TradingLab (Toàn bộ 109 Video)"
if sheet_name in wb.sheetnames:
    del wb[sheet_name]

ws = wb.create_sheet(title=sheet_name)

# Header
headers = [
    "STT", "Trạng thái Script", "Thời lượng", "Tiêu đề Video", 
    "Chủ đề cốt lõi", "Số dòng thoại", "Xem video YouTube", 
    "Đọc File Script cục bộ", "Trích đoạn thoại mở đầu"
]
ws.append(headers)

# Styling
header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

success_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
failed_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

for row_idx, v in enumerate(videos, start=2):
    status_str = "Đã có Full Script" if v["success"] else "Không có phụ đề"
    
    # Classify topic
    t_lower = v["title"].lower()
    topic = "Price Action"
    if "liquidity" in t_lower or "trap" in t_lower:
        topic = "Liquidity & Bẫy giá"
    elif "breakout" in t_lower or "fakeout" in t_lower:
        topic = "Breakout / Fakeout"
    elif "orderflow" in t_lower or "volume" in t_lower:
        topic = "Orderflow & Volume"
    elif "fibonacci" in t_lower:
        topic = "Fibonacci OTE"
    elif "indicator" in t_lower or "supertrend" in t_lower or "macd" in t_lower or "rsi" in t_lower:
        topic = "Chỉ báo kỹ thuật"
    elif "supply" in t_lower or "demand" in t_lower or "orderblock" in t_lower:
        topic = "Supply Demand & OB"
    elif "scalping" in t_lower or "day trade" in t_lower:
        topic = "Scalping & Day Trading"
        
    rel_file = os.path.join("tradinglab_transcripts", v["file_name"])
    
    ws.append([
        v["stt"],
        status_str,
        v["duration"],
        v["title"],
        topic,
        v["lines_count"],
        v["url"],
        rel_file,
        v["preview"]
    ])
    
    # Fill colors
    row_fill = success_fill if v["success"] else failed_fill
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center")
        if c in [1, 2, 3, 5, 6]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        if c == 2:
            cell.fill = row_fill

# Column widths
widths = [6, 18, 12, 45, 22, 14, 35, 45, 60]
for idx, w in enumerate(widths, 1):
    col_letter = openpyxl.utils.get_column_letter(idx)
    ws.column_dimensions[col_letter].width = w

wb.save(wb_path)
print("Đã lưu bảng tính MASTER_VIDEO_DATABASE.xlsx với Sheet mới: TradingLab (109 Video)!")

# 2. Generate TRADINGLAB_SCRIPTS_HUB.md
md_path = os.path.join(research_dir, "TRADINGLAB_SCRIPTS_HUB.md")

with open(md_path, "w", encoding="utf-8") as f_md:
    f_md.write("""# 🏛️ KHO TOÀN BỘ SCRIPT LỜI THOẠI KÊNH TRADINGLAB (109 VIDEOS)

> **Giới thiệu:** Toàn bộ lời thoại chi tiết (full transcript có timestamp) của 109 video trên kênh **TradingLab**.
> - **Thư mục script chi tiết:** Mỗi video được lưu riêng thành 1 file markdown độc lập tại `tradinglab_transcripts/` để dễ tra cứu, đọc không bị đơ giật máy.
> - **Bảng tính Excel tổng thể:** [MASTER_VIDEO_DATABASE.xlsx](file:///D:/victor/Trading/docs/research/MASTER_VIDEO_DATABASE.xlsx) (Xem sheet `TradingLab (Toàn bộ 109 Video)`).
> - **Trung tâm điều hành chính:** [MASTER_VIDEO_HUB.md](file:///D:/victor/Trading/docs/research/MASTER_VIDEO_HUB.md).

---

## 📑 BẢNG MỤC LỤC TOÀN BỘ 109 VIDEO TRADINGLAB

| STT | Trạng thái | Thời lượng | Chủ đề | Tiêu đề Video | Đọc Lời thoại (Local Script) | Link YouTube |
|:---:|:---:|:---:|---|---|:---:|:---:|
""")
    for v in videos:
        status_badge = "✅ Full Script" if v["success"] else "⚠️ Không sub"
        t_lower = v["title"].lower()
        topic = "Price Action"
        if "liquidity" in t_lower or "trap" in t_lower:
            topic = "Liquidity & Bẫy"
        elif "breakout" in t_lower or "fakeout" in t_lower:
            topic = "Breakout / Fakeout"
        elif "orderflow" in t_lower or "volume" in t_lower:
            topic = "Orderflow & Volume"
        elif "fibonacci" in t_lower:
            topic = "Fibonacci OTE"
        elif "indicator" in t_lower or "supertrend" in t_lower or "macd" in t_lower or "rsi" in t_lower:
            topic = "Chỉ báo kỹ thuật"
        elif "supply" in t_lower or "demand" in t_lower or "orderblock" in t_lower:
            topic = "Supply/Demand & OB"
        elif "scalping" in t_lower or "day trade" in t_lower:
            topic = "Scalping / Day Trade"
            
        rel_path = f"tradinglab_transcripts/{v['file_name']}"
        f_md.write(f"| **#{v['stt']}** | {status_badge} | {v['duration']} | `{topic}` | {v['title']} | [📖 Đọc Script]({rel_path}) | [Xem YouTube]({v['url']}) |\n")

print("Đã tạo TRADINGLAB_SCRIPTS_HUB.md thành công!")

# 3. Add link banner into MASTER_VIDEO_HUB.md and README.md
for hub_file in ["MASTER_VIDEO_HUB.md", "README.md"]:
    h_path = os.path.join(research_dir, hub_file)
    with open(h_path, "r", encoding="utf-8") as f_h:
        content = f_h.read()
    
    banner = """> [!TIP]
> **🏛️ ĐÃ LƯU TRỌN BỘ 109 SCRIPT KÊNH TRADINGLAB:**
> Toàn bộ lời thoại chi tiết kèm timestamp của **109 video kênh TradingLab** đã được tải về đầy đủ tại thư mục [tradinglab_transcripts/](file:///D:/victor/Trading/docs/research/tradinglab_transcripts).
> - Xem mục lục tra cứu nhanh: [TRADINGLAB_SCRIPTS_HUB.md](file:///D:/victor/Trading/docs/research/TRADINGLAB_SCRIPTS_HUB.md).
> - Xem trên Excel: [MASTER_VIDEO_DATABASE.xlsx](file:///D:/victor/Trading/docs/research/MASTER_VIDEO_DATABASE.xlsx) (Sheet: `TradingLab (Toàn bộ 109 Video)`).
"""
    if "TRADINGLAB_SCRIPTS_HUB.md" not in content:
        # insert right under line 8
        content = content.replace("---\n\n## 🚨 BẢNG HÀNG ĐỢI XEM", banner + "\n---\n\n## 🚨 BẢNG HÀNG ĐỢI XEM")
        with open(h_path, "w", encoding="utf-8") as f_h:
            f_h.write(content)
        print(f"Đã chèn banner kho TradingLab vào {hub_file}!")

